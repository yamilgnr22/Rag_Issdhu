import io
import logging
import mimetypes
import os
import re
import sys
import tempfile
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pdfplumber
from bs4 import BeautifulSoup
from docx import Document as DocxDocument

from app.config import Settings


logger = logging.getLogger(__name__)


@dataclass
class ExtractionBlock:
    block_type: str
    text: str
    page: int
    section_path: list[str]


@dataclass
class ExtractionResult:
    blocks: list[ExtractionBlock]
    backend: str
    review_required: bool
    review_reason: str | None = None
    quality: dict[str, float] | None = None
    failed_pages: list[int] | None = None

    @property
    def text_length(self) -> int:
        return sum(len(block.text) for block in self.blocks)


WORD_PATTERN = re.compile(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]{2,}")
ACCENTED = set("áéíóúüñÁÉÍÓÚÜÑ")
# docling-parse reporta las paginas que no pudo procesar por stderr nativo y
# devuelve el resto del documento como si la conversion hubiera sido completa:
#   "Stage preprocess failed for run 1, pages [223]: std::bad_alloc"
FAILED_PAGES_PATTERN = re.compile(r"failed[^\n]*?pages?\s*\[([\d,\s]+)\]", re.IGNORECASE)


@contextmanager
def capture_native_stderr():
    """Captura stderr a nivel de descriptor.

    contextlib.redirect_stderr solo desvia sys.stderr, y estos mensajes los
    escribe una libreria nativa directamente al fd 2.
    """
    try:
        target_fd = sys.stderr.fileno()
    except Exception:
        yield None  # sin fd real (p.ej. stderr capturado): no se puede interceptar
        return

    saved_fd = os.dup(target_fd)
    tmp = tempfile.TemporaryFile(mode="w+b")
    try:
        sys.stderr.flush()
        os.dup2(tmp.fileno(), target_fd)
        yield tmp
    finally:
        try:
            sys.stderr.flush()
        except Exception:
            pass
        os.dup2(saved_fd, target_fd)
        os.close(saved_fd)
        try:
            tmp.seek(0)
            captured = tmp.read().decode("utf-8", errors="replace")
            if captured.strip():
                sys.stderr.write(captured)  # no tragarse la salida original
        except Exception:
            pass
        tmp.close()


def parse_failed_pages(stderr_text: str) -> list[int]:
    pages: set[int] = set()
    for match in FAILED_PAGES_PATTERN.finditer(stderr_text or ""):
        for raw in match.group(1).split(","):
            raw = raw.strip()
            if raw.isdigit():
                pages.add(int(raw))
    return sorted(pages)


def assess_text_quality(text: str) -> dict[str, float]:
    """Mide dos sintomas de OCR degradado que se detectan sin conocer el documento.

    - glued_ratio: proporcion de "palabras" de 18 o mas letras. Cuando el OCR
      pierde los espacios aparecen tokens como "Delassiglasyconceptos".
    - accent_ratio: proporcion de palabras con tilde o enie. Un OCR que corrio
      sin el idioma correcto se come los diacriticos.
    """
    words = WORD_PATTERN.findall(text or "")
    total = len(words)
    if not total:
        return {"words": 0.0, "glued_ratio": 0.0, "accent_ratio": 0.0}
    glued = sum(1 for word in words if len(word) >= 18)
    accented = sum(1 for word in words if any(char in ACCENTED for char in word))
    return {
        "words": float(total),
        "glued_ratio": round(glued / total, 4),
        "accent_ratio": round(accented / total, 4),
    }


class DoclingExtractor:
    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings

    def _ocr_languages(self) -> list[str]:
        raw = getattr(self.settings, "ocr_language", None) or "spa+eng"
        return [part.strip() for part in re.split(r"[+,\s]+", raw) if part.strip()]

    def _has_text_layer(self, content: bytes) -> bool:
        """True si el PDF trae texto embebido. Un PDF escaneado devuelve False y
        necesita OCR con el idioma correcto: sin idioma, el motor pierde tildes y
        pega las palabras."""
        try:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                pages = pdf.pages
                if not pages:
                    return False
                sample = [pages[index] for index in range(0, len(pages), max(1, len(pages) // 8))]
                return any(len((page.extract_text() or "").strip()) > 80 for page in sample)
        except Exception:
            return True

    def _pdf_page_count(self, content: bytes) -> int:
        try:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                return len(pdf.pages)
        except Exception:
            return 0

    def _convert_document(self, converter, path: Path, *, total_pages: int) -> tuple[str, list[int]]:
        """Convierte el documento, por tramos si es grande.

        Devuelve (markdown, paginas_perdidas). Trocear acota el pico de memoria
        de docling; sin esto, los documentos largos pierden paginas en silencio.
        """
        batch_size = int(getattr(self.settings, "extraction_batch_pages", 0) or 0)
        threshold = int(getattr(self.settings, "extraction_batch_threshold_pages", 0) or 0)
        use_batches = bool(batch_size and total_pages and threshold and total_pages > threshold)

        ranges: list[tuple[int, int]]
        if use_batches:
            ranges = [
                (start, min(start + batch_size - 1, total_pages))
                for start in range(1, total_pages + 1, batch_size)
            ]
            logger.info(
                "extraction_batched file=%s pages=%d batches=%d size=%d",
                path.name,
                total_pages,
                len(ranges),
                batch_size,
            )
        else:
            ranges = [(1, total_pages)] if total_pages else [(1, 10**9)]

        parts: list[str] = []
        failed: list[int] = []
        for index, (start, end) in enumerate(ranges, start=1):
            with capture_native_stderr() as buffer:
                if use_batches:
                    result = converter.convert(str(path), page_range=(start, end))
                else:
                    result = converter.convert(str(path))
                captured = ""
                if buffer is not None:
                    buffer.flush()
                    buffer.seek(0)
                    captured = buffer.read().decode("utf-8", errors="replace")
            lost = parse_failed_pages(captured)
            if lost:
                logger.error(
                    "extraction_pages_lost file=%s batch=%d/%d range=%s-%s pages=%s",
                    path.name,
                    index,
                    len(ranges),
                    start,
                    end,
                    lost[:20],
                )
                failed.extend(lost)
            markdown = result.document.export_to_markdown().strip()
            if markdown:
                parts.append(markdown)

        return "\n\n".join(parts).strip(), sorted(set(failed))

    def _build_converter(self, *, needs_ocr: bool):
        from docling.document_converter import DocumentConverter, PdfFormatOption
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import PdfPipelineOptions, TesseractCliOcrOptions

        pipeline_options = PdfPipelineOptions()
        pipeline_options.do_ocr = True
        pipeline_options.ocr_options = TesseractCliOcrOptions(
            lang=self._ocr_languages(),
            force_full_page_ocr=needs_ocr,
        )
        return DocumentConverter(
            format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options)}
        )

    def extract(self, filename: str, content: bytes) -> ExtractionResult | None:
        try:
            from docling.document_converter import DocumentConverter
        except Exception:
            return None

        suffix = Path(filename).suffix.lower()
        if suffix not in {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm", ".png", ".jpg", ".jpeg"}:
            return None

        temp_dir = Path(".data") / "docling_tmp"
        temp_dir.mkdir(parents=True, exist_ok=True)
        temp_path = temp_dir / f"docling_{Path(filename).name}"
        temp_path.write_bytes(content)

        is_pdf = suffix == ".pdf"
        needs_ocr = is_pdf and not self._has_text_layer(content)

        try:
            try:
                converter = self._build_converter(needs_ocr=needs_ocr) if is_pdf else DocumentConverter()
            except Exception:
                converter = DocumentConverter()

            markdown, failed_pages = self._convert_document(
                converter,
                temp_path,
                total_pages=self._pdf_page_count(content) if is_pdf else 0,
            )
            if failed_pages:
                logger.error(
                    "extraction_pages_lost_total file=%s pages=%d %s "
                    "(docling descarto estas paginas y devolvio el resto como conversion completa)",
                    filename,
                    len(failed_pages),
                    failed_pages[:20],
                )
            if not markdown:
                return None
            review_required = len(markdown) < 50
            reason = "docling_extracted_too_little_text" if review_required else None
            if failed_pages:
                review_required = True
                lost = f"pages_lost_in_extraction:{len(failed_pages)}"
                reason = f"{reason};{lost}" if reason else lost
            backend = "docling"
            if needs_ocr:
                backend = f"docling+ocr[{'+'.join(self._ocr_languages())}]"
            return ExtractionResult(
                blocks=[ExtractionBlock("docling_markdown", markdown, 1, ["Document"])],
                backend=backend,
                review_required=review_required,
                review_reason=reason,
                failed_pages=failed_pages or None,
            )
        except Exception:
            return None
        finally:
            try:
                temp_path.unlink(missing_ok=True)
            except Exception:
                pass


class FallbackExtractor:
    MIN_TEXT_LENGTH = 50

    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def extract(self, filename: str, content: bytes, mime_type: str | None = None) -> ExtractionResult:
        suffix = Path(filename).suffix.lower()
        mime_type = mime_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"

        if suffix in {".txt", ".md", ".json"} or mime_type.startswith("text/"):
            text = self._decode_text(content)
            return self._finalize([ExtractionBlock("text", text, 1, ["Document"])], "plain_text")

        if suffix == ".pdf" or mime_type == "application/pdf":
            return self._extract_pdf(content)

        if suffix == ".docx":
            return self._extract_docx(content)

        if suffix in {".html", ".htm"} or mime_type == "text/html":
            return self._extract_html(content)

        if suffix == ".xlsx":
            return self._extract_xlsx(content)

        return ExtractionResult([], "unsupported", True, f"unsupported_format:{suffix or mime_type}")

    def _decode_text(self, content: bytes) -> str:
        for encoding in ("utf-8", "cp1252", "latin-1"):
            try:
                return content.decode(encoding).strip()
            except UnicodeDecodeError:
                continue
        return ""

    def _extract_pdf(self, content: bytes) -> ExtractionResult:
        blocks: list[ExtractionBlock] = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    blocks.append(
                        ExtractionBlock(
                            "page",
                            re.sub(r"\s+\n", "\n", text),
                            page_number,
                            [f"Page {page_number}"],
                        )
                    )
        if not blocks:
            return ExtractionResult([], "pdfplumber", True, "pdf_no_text_detected_manual_review")
        return self._finalize(blocks, "pdfplumber")

    def _extract_docx(self, content: bytes) -> ExtractionResult:
        doc = DocxDocument(io.BytesIO(content))
        blocks: list[ExtractionBlock] = []
        current_heading = "Document"
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            if para.style and para.style.name.lower().startswith("heading"):
                current_heading = text
                blocks.append(ExtractionBlock("heading", text, 1, [current_heading]))
            else:
                blocks.append(ExtractionBlock("paragraph", text, 1, [current_heading]))
        for table_index, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    rows.append("\t".join(values))
            if rows:
                blocks.append(
                    ExtractionBlock("table", "\n".join(rows), 1, [current_heading, f"Table {table_index}"])
                )
        return self._finalize(blocks, "python-docx")

    def _extract_html(self, content: bytes) -> ExtractionResult:
        soup = BeautifulSoup(content, "html.parser")
        blocks: list[ExtractionBlock] = []
        current_heading = "Document"
        for tag in soup.find_all(["h1", "h2", "h3", "p", "li", "table"]):
            text = tag.get_text(" ", strip=True)
            if not text:
                continue
            if tag.name in {"h1", "h2", "h3"}:
                current_heading = text
                blocks.append(ExtractionBlock("heading", text, 1, [current_heading]))
            elif tag.name == "table":
                blocks.append(ExtractionBlock("table", text, 1, [current_heading, "Table"]))
            else:
                block_type = "paragraph" if tag.name == "p" else "list_item"
                blocks.append(ExtractionBlock(block_type, text, 1, [current_heading]))
        return self._finalize(blocks, "beautifulsoup4")

    def _extract_xlsx(self, content: bytes) -> ExtractionResult:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        blocks: list[ExtractionBlock] = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    rows.append("\t".join(values))
            if rows:
                blocks.append(ExtractionBlock("table", "\n".join(rows), 1, [sheet.title]))
        return self._finalize(blocks, "openpyxl")

    def _finalize(self, blocks: list[ExtractionBlock], backend: str) -> ExtractionResult:
        cleaned = [block for block in blocks if block.text.strip()]
        text_length = sum(len(block.text) for block in cleaned)
        review_required = text_length < self.MIN_TEXT_LENGTH
        return ExtractionResult(
            blocks=cleaned,
            backend=backend,
            review_required=review_required,
            review_reason="extracted_text_below_threshold" if review_required else None,
        )


class ExtractionPipeline:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self.docling = DoclingExtractor(settings)
        self.fallback = FallbackExtractor(settings)

    def extract(self, filename: str, content: bytes) -> ExtractionResult:
        docling_result = self.docling.extract(filename, content)
        result = docling_result if (docling_result and docling_result.blocks) else self.fallback.extract(filename, content)
        return self._apply_quality_gate(result, filename=filename, content=content)

    def _reference_chars_per_page(self, content: bytes) -> tuple[float, int]:
        """Caracteres por pagina segun pdfplumber, como referencia independiente
        de lo que haya extraido docling. Devuelve (chars_por_pagina, paginas)."""
        try:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                pages = pdf.pages
                total = len(pages)
                if not total:
                    return 0.0, 0
                sample_index = list(range(0, total, max(1, total // 12)))
                sampled = sum(len((pages[i].extract_text() or "")) for i in sample_index)
                return sampled / len(sample_index), total
        except Exception:
            return 0.0, 0

    def _apply_quality_gate(
        self,
        result: ExtractionResult,
        *,
        filename: str = "",
        content: bytes | None = None,
    ) -> ExtractionResult:
        """Marca para revision manual una extraccion legible por el parser pero
        inservible para recuperacion, o incompleta."""
        if not getattr(self.settings, "extraction_quality_gate_enabled", True):
            return result

        text = " ".join(block.text for block in result.blocks)
        quality = assess_text_quality(text)

        # Integridad: solo aplicable a PDFs con capa de texto, que es donde
        # pdfplumber sirve de referencia independiente.
        if content and Path(filename).suffix.lower() == ".pdf":
            reference, pages = self._reference_chars_per_page(content)
            if reference > 200 and pages:
                extracted_per_page = len(text) / pages
                quality["pages"] = float(pages)
                quality["chars_per_page"] = round(extracted_per_page, 1)
                quality["reference_chars_per_page"] = round(reference, 1)
                quality["coverage_ratio"] = round(extracted_per_page / reference, 3)

        result.quality = quality

        # En textos cortos las proporciones son ruido; el gate no opina.
        if quality["words"] < float(self.settings.extraction_quality_min_words):
            return result

        reasons: list[str] = []
        if quality["glued_ratio"] > float(self.settings.extraction_max_glued_ratio):
            reasons.append(f"glued_tokens={quality['glued_ratio']:.3f}")
        if quality["accent_ratio"] < float(self.settings.extraction_min_accent_ratio):
            reasons.append(f"missing_accents={quality['accent_ratio']:.3f}")
        coverage = quality.get("coverage_ratio")
        if coverage is not None and coverage < float(self.settings.extraction_min_coverage_ratio):
            reasons.append(f"incomplete_extraction={coverage:.2f}")
        if result.failed_pages:
            reasons.append(f"pages_lost={len(result.failed_pages)}")

        if reasons:
            logger.warning(
                "extraction_quality_gate backend=%s words=%d %s",
                result.backend,
                int(quality["words"]),
                " ".join(reasons),
            )
            result.review_required = True
            existing = result.review_reason
            gate_reason = "low_extraction_quality:" + ",".join(reasons)
            result.review_reason = f"{existing};{gate_reason}" if existing else gate_reason
        return result

